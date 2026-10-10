import csv, io, re, zipfile
from pathlib import PurePosixPath
from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import load_workbook
from PIL import Image
from .models import Module,Category,Question,Option
MAX_UNPACKED=50*1024*1024
HEADERS=["question_code","module_code","category_code","category_title","question_type","question_text","option_a","option_b","option_c","option_d","option_e","option_f","correct_options","marks","difficulty","explanation","image_filename","image_alt_text","status"]
ALLOWED_IMAGES={".jpg",".jpeg",".png",".webp"}; MAX_IMAGE=5*1024*1024

def clean(v): return str(v).strip() if v is not None else ""
def read_rows(name,data):
    if name.lower().endswith(".csv"):
        text=data.decode("utf-8-sig"); return list(csv.DictReader(io.StringIO(text)))
    if name.lower().endswith(".xlsx"):
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); ws=wb["Questions"] if "Questions" in wb.sheetnames else wb.active
        values=ws.iter_rows(values_only=True); headers=[clean(x) for x in next(values)]
        return [dict(zip(headers,row)) for row in values if any(v is not None and clean(v) for v in row)]
    raise ValueError("Use CSV, XLSX, or ZIP containing questions.csv/questions.xlsx.")

def unpack(upload):
    raw=upload.read(); name=upload.name; assets={}
    if not name.lower().endswith(".zip"): return name,raw,assets
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        safe=[]; total_unpacked=0
        for info in z.infolist():
            p=PurePosixPath(info.filename)
            if p.is_absolute() or ".." in p.parts: raise ValueError("ZIP contains an unsafe path.")
            if not info.is_dir():
                total_unpacked += info.file_size
                if total_unpacked > MAX_UNPACKED: raise ValueError("ZIP expands beyond the 50 MB safety limit.")
                safe.append(info)
        sheets=[x for x in safe if PurePosixPath(x.filename).name.lower() in {"questions.csv","questions.xlsx"}]
        if len(sheets)!=1: raise ValueError("ZIP must contain exactly one questions.csv or questions.xlsx file.")
        sheet=sheets[0]; sheet_data=z.read(sheet)
        for info in safe:
            ext=PurePosixPath(info.filename).suffix.lower()
            if ext in ALLOWED_IMAGES:
                basename=PurePosixPath(info.filename).name
                if basename in assets: raise ValueError(f"ZIP contains duplicate image basename: {basename}.")
                data=z.read(info)
                if len(data)>MAX_IMAGE: raise ValueError(f"Image {info.filename} exceeds 5 MB.")
                try:
                    im=Image.open(io.BytesIO(data)); im.verify()
                except Exception: raise ValueError(f"Image {info.filename} is invalid.")
                assets[basename]=data
        return PurePosixPath(sheet.filename).name,sheet_data,assets

def validate_rows(rows,assets):
    errors=[]; parsed=[]; seen=set()
    for n,row in enumerate(rows,2):
        r={h:clean(row.get(h,"")) for h in HEADERS}; es=[]
        code=r["question_code"]
        if not code: es.append("question_code is required")
        if code in seen: es.append("duplicate question_code in file")
        seen.add(code)
        if not re.fullmatch(r"[A-Za-z0-9_.-]{1,60}",code): es.append("question_code must use letters, numbers, ., _, or -")
        if not r["module_code"]: es.append("module_code is required")
        if not r["category_code"]: es.append("category_code is required")
        if r["question_type"].upper() not in {"SINGLE","MULTIPLE"}: es.append("question_type must be SINGLE or MULTIPLE")
        if not r["question_text"]: es.append("question_text is required")
        opts={chr(65+i):r[f"option_{chr(97+i)}"] for i in range(6)}; opts={k:v for k,v in opts.items() if v}
        if len(opts)<2: es.append("at least two options are required")
        correct={x.strip().upper() for x in r["correct_options"].split(",") if x.strip()}
        if not correct or not correct<=set(opts): es.append("correct_options must name existing options")
        if r["question_type"].upper()=="SINGLE" and len(correct)!=1: es.append("SINGLE requires exactly one correct option")
        if r["question_type"].upper()=="MULTIPLE" and len(correct)<2: es.append("MULTIPLE requires at least two correct options")
        try:
            marks=float(r["marks"] or 1)
            if marks<=0: raise ValueError
        except: es.append("marks must be greater than zero"); marks=1
        diff=(r["difficulty"] or "MEDIUM").upper()
        if diff not in {"EASY","MEDIUM","HARD"}: es.append("difficulty must be EASY, MEDIUM, or HARD")
        status=(r["status"] or "DRAFT").upper()
        if status not in {"DRAFT","PUBLISHED","RETIRED"}: es.append("invalid status")
        image=r["image_filename"]
        if image and image not in assets: es.append("image_filename was not found in ZIP")
        if image and not r["image_alt_text"]: es.append("image_alt_text is required when an image is used")
        if es: errors.append({"row":n,"question_code":code,"errors":es})
        parsed.append((r,opts,correct,marks,diff,status))
    return parsed,errors

@transaction.atomic
def import_questions(upload,user,dry_run=False):
    name,data,assets=unpack(upload); rows=read_rows(name,data); parsed,errors=validate_rows(rows,assets)
    if errors: return {"ok":False,"rows":len(rows),"imported":0,"errors":errors}
    if dry_run: return {"ok":True,"rows":len(rows),"imported":0,"errors":[]}
    imported=0
    for r,opts,correct,marks,diff,status in parsed:
        module,_=Module.objects.get_or_create(code=r["module_code"].upper(),defaults={"title":r["module_code"].replace("_"," ").title()})
        category,_=Category.objects.get_or_create(module=module,code=r["category_code"].upper(),defaults={"title":r["category_title"] or r["category_code"].replace("_"," ").title()})
        q,created=Question.objects.update_or_create(code=r["question_code"],defaults={"module":module,"category":category,"question_type":r["question_type"].upper(),"stem":r["question_text"],"explanation":r["explanation"],"difficulty":diff,"marks":marks,"status":status,"image_alt_text":r["image_alt_text"],"created_by":user})
        q.options.all().delete(); Option.objects.bulk_create([Option(question=q,key=k,text=v,is_correct=k in correct) for k,v in opts.items()])
        if r["image_filename"]:
            ext=PurePosixPath(r["image_filename"]).suffix.lower(); q.image.save(f"{q.code}{ext}",ContentFile(assets[r["image_filename"]]),save=True)
        imported+=1
    return {"ok":True,"rows":len(rows),"imported":imported,"errors":[]}


# The default instructor workflow deliberately uses a small, human-readable sheet.
# The original 19-column importer above remains for backwards compatibility and tests.
SIMPLE_HEADERS = [
    "module", "category", "question", "option_a", "option_b", "option_c",
    "option_d", "correct_answer", "image", "image_description",
]

def _normalise_header(value):
    return re.sub(r"[^a-z0-9]+", "_", clean(value).lower()).strip("_")

def read_simple_rows(name, data):
    if name.lower().endswith(".csv"):
        reader = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        values = list(reader)
    elif name.lower().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook["Questions"] if "Questions" in workbook.sheetnames else workbook.active
        values = list(sheet.iter_rows(values_only=True))
    else:
        raise ValueError("Use the supplied XLSX template or a CSV file.")
    if not values:
        raise ValueError("The spreadsheet is empty.")
    headers = [_normalise_header(v) for v in values[0]]
    aliases = {
        "correct": "correct_answer", "correct_answers": "correct_answer", "answer": "correct_answer",
        "picture": "image", "picture_filename": "image", "image_filename": "image", "diagram": "image",
        "alt_text": "image_description", "image_alt_text": "image_description", "description": "image_description",
        "question_text": "question", "stem": "question",
        "topic": "category", "subtopic": "category", "section": "category", "sections": "category",
        "chapter": "category", "syllabus_section": "category",
    }
    headers = [aliases.get(h, h) for h in headers]
    # Note: 'module' is no longer required in the spreadsheet since exams are module-specific.
    missing = [h for h in ("question", "option_a", "option_b", "correct_answer") if h not in headers]
    if missing:
        raise ValueError("Missing required column(s): " + ", ".join(missing) + ". Download and use the simple template.")
    rows = []
    for values_row in values[1:]:
        row = {h: clean(v) for h, v in zip(headers, values_row)}
        if any(row.values()):
            rows.append(row)
    if not rows:
        raise ValueError("The spreadsheet has no question rows.")
    return rows

def _safe_code(value, fallback):
    code = re.sub(r"[^A-Z0-9]+", "_", clean(value).upper()).strip("_")
    return (code or fallback)[:50]

def _read_image_uploads(image_uploads):
    assets_by_name = {}
    assets_by_stem = {}
    errors = []
    for upload in image_uploads:
        path = PurePosixPath(upload.name)
        basename = path.name
        key_name = basename.lower()
        key_stem = path.stem.lower()
        if key_name in assets_by_name:
            errors.append(f"Two selected pictures have the same filename: {basename}")
            continue
        data = upload.read()
        if len(data) > MAX_IMAGE:
            errors.append(f"{basename} exceeds 5 MB")
            continue
        try:
            image = Image.open(io.BytesIO(data))
            image.verify()
            if image.format not in {"PNG", "JPEG", "WEBP"}:
                raise ValueError
        except (OSError, SyntaxError, ValueError):
            errors.append(f"{basename} is not a valid PNG, JPEG, or WebP image")
            continue
        assets_by_name[key_name] = (basename, data)
        # Format-agnostic indexing: allow looking up 'clouds' when 'clouds.png' was uploaded
        if key_stem not in assets_by_stem:
            assets_by_stem[key_stem] = (basename, data)
    return assets_by_name, assets_by_stem, errors

def validate_simple_rows(rows, assets_by_name, assets_by_stem, default_module=None):
    parsed, errors = [], []
    for row_number, source in enumerate(rows, 2):
        row = {h: clean(source.get(h, "")) for h in SIMPLE_HEADERS}
        row_errors = []
        module_title = default_module or row["module"] or "General Aviation"
        category_title = row["category"] or "General"
        if not module_title:
            row_errors.append("Module / Subject is required")
        if not row["question"]:
            row_errors.append("Question is required")
        options = {letter: row[f"option_{letter.lower()}"] for letter in "ABCD"}
        options = {letter: text for letter, text in options.items() if text}
        if len(options) < 2:
            row_errors.append("Enter at least Option A and Option B")
        correct = {x.strip().upper() for x in re.split(r"[,;/ ]+", row["correct_answer"]) if x.strip()}
        if not correct:
            row_errors.append("Correct answer is required (for example A or A,C)")
        elif not correct <= set(options):
            row_errors.append("Correct answer must name an option that exists")
        
        # Format-agnostic image resolution
        raw_image = row["image"]
        resolved_image_name = ""
        if raw_image:
            img_path = PurePosixPath(raw_image)
            lower_name = img_path.name.lower()
            lower_stem = img_path.stem.lower()
            match = assets_by_stem.get(lower_stem) or assets_by_name.get(lower_name)
            if not match:
                row_errors.append(f"Picture '{raw_image}' was not selected in upload")
            else:
                resolved_image_name = match[0]
            if not row["image_description"]:
                row_errors.append("Image description is required when a picture is used")

        if row_errors:
            errors.append({"row": row_number, "question_code": "", "errors": row_errors})
        parsed.append({
            "module_title": module_title,
            "module_code": _safe_code(module_title, "MODULE"),
            "category_title": category_title,
            "category_code": _safe_code(category_title, "GENERAL"),
            "stem": row["question"], "options": options, "correct": correct,
            "question_type": Question.MULTIPLE if len(correct) > 1 else Question.SINGLE,
            "image_name": resolved_image_name, "image_alt_text": row["image_description"],
        })
    return parsed, errors

@transaction.atomic
def import_simple_questions(spreadsheet, image_uploads, user, default_module=None):
    raw = spreadsheet.read()
    rows = read_simple_rows(spreadsheet.name, raw)
    assets_by_name, assets_by_stem, asset_errors = _read_image_uploads(image_uploads)
    parsed, errors = validate_simple_rows(rows, assets_by_name, assets_by_stem, default_module=default_module)
    if asset_errors:
        errors.insert(0, {"row": "Pictures", "question_code": "", "errors": asset_errors})
    if errors:
        return {"ok": False, "rows": len(rows), "imported": 0, "errors": errors, "unused_images": []}
    imported = 0
    imported_ids = []
    used_images = set()
    for item in parsed:
        module, _ = Module.objects.get_or_create(
            code=item["module_code"], defaults={"title": item["module_title"]}
        )
        category, _ = Category.objects.get_or_create(
            module=module, code=item["category_code"],
            defaults={"title": item["category_title"]},
        )
        import uuid
        code = f"{item['module_code'][:30]}-{uuid.uuid4().hex[:10].upper()}"
        question = Question.objects.create(
            code=code, module=module, category=category,
            question_type=item["question_type"], stem=item["stem"],
            marks=1, difficulty="MEDIUM", status=Question.DRAFT,
            image_alt_text=item["image_alt_text"], created_by=user,
        )
        Option.objects.bulk_create([
            Option(question=question, key=key, text=text, is_correct=key in item["correct"])
            for key, text in item["options"].items()
        ])
        if item["image_name"]:
            original_name, data = assets_by_name[item["image_name"].lower()]
            extension = PurePosixPath(original_name).suffix.lower()
            question.image.save(f"{question.code}{extension}", ContentFile(data), save=True)
            used_images.add(item["image_name"].lower())
        imported += 1
        imported_ids.append(question.id)
    unused = sorted(name for name, (original, _) in assets_by_name.items() if name not in used_images)
    return {"ok": True, "rows": len(rows), "imported": imported, "imported_ids": imported_ids, "errors": [], "unused_images": unused}
