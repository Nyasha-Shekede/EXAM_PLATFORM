import io, zipfile
from datetime import timedelta
from PIL import Image
from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase,override_settings
from django.urls import reverse
from django.utils import timezone
from .importers import import_questions, import_simple_questions
from .models import *
from .services import start_attempt,save_response,submit_attempt,verify_audit_chain,validate_question

class Base(TestCase):
    def setUp(self):
        U=get_user_model(); self.user=U.objects.create_user("C001","candidate@example.test","Strong-pass-473",first_name="Avi",last_name="Student"); self.other=U.objects.create_user("C002",password="Strong-pass-474"); self.staff=U.objects.create_superuser("admin","admin@example.test","Strong-pass-475")
        self.module=Module.objects.create(code="AIR_LAW",title="Air Law")
        self.cat=Category.objects.create(module=self.module,code="RULES",title="Rules")
        for i in range(5):
            q=Question.objects.create(code=f"Q{i+1}",module=self.module,category=self.cat,question_type=Question.SINGLE,stem=f"Question {i+1}",status=Question.PUBLISHED,created_by=self.staff)
            Option.objects.create(question=q,key="A",text="Correct",is_correct=True); Option.objects.create(question=q,key="B",text="Wrong")
        self.exam=Exam.objects.create(code="AIR-EXAM",title="Air Law Exam",module=self.module,duration_minutes=30,question_count=3,pass_mark=75,max_attempts=2,status=Exam.PUBLISHED,show_answers_after=True)
        Assignment.objects.create(exam=self.exam,candidate=self.user)

class ModelAndScoringTests(Base):
    def test_question_validation_rules(self):
        q=Question.objects.create(code="M1",module=self.module,category=self.cat,question_type=Question.MULTIPLE,stem="Pick",status=Question.DRAFT)
        Option.objects.create(question=q,key="A",text="A",is_correct=True); Option.objects.create(question=q,key="B",text="B")
        self.assertIn("at least two correct",validate_question(q)[0])
    def test_start_creates_snapshot_without_answer_leak_in_page(self):
        a=start_attempt(self.exam,self.user); self.assertEqual(a.attempt_questions.count(),3); self.assertEqual(AttemptOption.objects.filter(attempt_question__attempt=a).count(),6)
        self.client.login(username="C001",password="Strong-pass-473"); html=self.client.get(reverse("question",args=[a.id,1])).content.decode(); self.assertNotIn("is_correct",html); self.assertNotIn("Correct answer",html)
    def test_multi_choice_exact_match_no_partial_credit(self):
        q=Question.objects.create(code="MULTI",module=self.module,category=self.cat,question_type=Question.MULTIPLE,stem="Choose",marks=2,status=Question.PUBLISHED)
        for key,ok in [("A",True),("B",True),("C",False)]: Option.objects.create(question=q,key=key,text=key,is_correct=ok)
        Question.objects.filter(code__startswith="Q").update(status=Question.RETIRED)
        self.exam.question_count=1; self.exam.save(); a=start_attempt(self.exam,self.user); aq=a.attempt_questions.get(); correct=list(aq.snapshot_options.filter(is_correct=True).values_list("display_key",flat=True))
        save_response(a,aq,correct[:1],self.user); submit_attempt(a,actor=self.user); a.refresh_from_db(); self.assertEqual(a.score,0); self.assertFalse(a.passed)
    def test_full_correct_score_passes(self):
        self.exam.question_count=1; self.exam.save(); a=start_attempt(self.exam,self.user); aq=a.attempt_questions.get(); keys=list(aq.snapshot_options.filter(is_correct=True).values_list("display_key",flat=True)); save_response(a,aq,keys,self.user); submit_attempt(a,actor=self.user); a.refresh_from_db(); self.assertEqual(a.percentage,100); self.assertTrue(a.passed)
    def test_expired_attempt_rejects_save_and_is_marked(self):
        self.exam.question_count=1; self.exam.save(); a=start_attempt(self.exam,self.user); Attempt.objects.filter(pk=a.pk).update(expires_at=timezone.now()-timedelta(seconds=1)); a.refresh_from_db()
        with self.assertRaises(ValidationError): save_response(a,a.attempt_questions.get(),["A"],self.user)
        a.refresh_from_db(); self.assertEqual(a.status,Attempt.EXPIRED)
    def test_resume_returns_same_attempt(self):
        a=start_attempt(self.exam,self.user); b=start_attempt(self.exam,self.user); self.assertEqual(a.id,b.id); self.assertEqual(Attempt.objects.count(),1)
    def test_attempt_limit(self):
        self.exam.question_count=1; self.exam.max_attempts=1; self.exam.save(); a=start_attempt(self.exam,self.user); submit_attempt(a)
        with self.assertRaises(ValidationError): start_attempt(self.exam,self.user)
    def test_unassigned_candidate_blocked(self):
        with self.assertRaises(ValidationError): start_attempt(self.exam,self.other)
    def test_snapshot_survives_source_edit(self):
        self.exam.question_count=1; self.exam.save(); a=start_attempt(self.exam,self.user); aq=a.attempt_questions.get(); old=aq.stem; q=aq.source_question; q.stem="Changed later"; q.save(); aq.refresh_from_db(); self.assertEqual(aq.stem,old)
    def test_audit_chain_links_events(self):
        a=start_attempt(self.exam,self.user); submit_attempt(a,actor=self.user); self.assertEqual(verify_audit_chain(),(True,None)); self.assertEqual(AuditEvent.objects.count(),2)
    def test_audit_chain_detects_payload_tampering(self):
        a=start_attempt(self.exam,self.user); submit_attempt(a,actor=self.user); first=AuditEvent.objects.first(); AuditEvent.objects.filter(pk=first.pk).update(details={"altered":True}); self.assertEqual(verify_audit_chain(),(False,first.id))

class WebTests(Base):
    def setUp(self): super().setUp(); self.client.login(username="C001",password="Strong-pass-473")
    def test_dashboard_and_begin(self):
        self.assertContains(self.client.get(reverse("dashboard")),"Air Law Exam"); r=self.client.post(reverse("begin",args=[self.exam.id])); self.assertEqual(r.status_code,302); self.assertIn("/questions/1/",r.url)
    def test_candidate_cannot_open_someone_elses_attempt(self):
        Assignment.objects.create(exam=self.exam,candidate=self.other); a=start_attempt(self.exam,self.other); self.assertEqual(self.client.get(reverse("question",args=[a.id,1])).status_code,404)
    def test_answer_autosave_endpoint(self):
        a=start_attempt(self.exam,self.user); aq=a.attempt_questions.get(position=1); key=aq.snapshot_options.first().display_key
        r=self.client.post(reverse("answer",args=[a.id,1]),{"selected":key},HTTP_X_REQUESTED_WITH="XMLHttpRequest"); self.assertEqual(r.status_code,200); self.assertEqual(aq.response.selected_keys,[key])
    def test_pdf_result_generated(self):
        self.exam.question_count=1; self.exam.save(); a=start_attempt(self.exam,self.user); submit_attempt(a); r=self.client.get(reverse("result_pdf",args=[a.id])); self.assertEqual(r.status_code,200); self.assertEqual(r["Content-Type"],"application/pdf"); self.assertTrue(r.content.startswith(b"%PDF"))
    def test_review_respects_setting(self):
        self.exam.question_count=1; self.exam.show_answers_after=False; self.exam.save(); a=start_attempt(self.exam,self.user); submit_attempt(a); self.assertEqual(self.client.get(reverse("review",args=[a.id])).status_code,404)
    def test_attempt_image_requires_owner(self):
        from django.core.files.base import ContentFile
        q=Question.objects.get(code="Q1"); im=Image.new("RGB",(10,10),(1,2,3)); b=io.BytesIO(); im.save(b,"PNG"); q.image.save("owned.png",ContentFile(b.getvalue()),save=True)
        self.exam.question_count=5; self.exam.save(); a=start_attempt(self.exam,self.user); aq=a.attempt_questions.get(source_question=q); url=reverse("attempt_image",args=[a.id,aq.position]); self.assertEqual(self.client.get(url).status_code,200)
        self.client.logout(); self.client.login(username="C002",password="Strong-pass-474"); self.assertEqual(self.client.get(url).status_code,404)
    def test_email_login_backend(self):
        self.client.logout(); self.assertTrue(self.client.login(username="candidate@example.test",password="Strong-pass-473"))

class ImportTests(Base):
    def csv_upload(self,text): return SimpleUploadedFile("questions.csv",text.encode(),content_type="text/csv")
    def test_valid_csv_import(self):
        csv="question_code,module_code,category_code,category_title,question_type,question_text,option_a,option_b,option_c,option_d,option_e,option_f,correct_options,marks,difficulty,explanation,image_filename,image_alt_text,status\nNEW-1,MET,METAR,Weather,SINGLE,What?,Yes,No,,,,,A,1,EASY,Because,,,PUBLISHED\n"
        result=import_questions(self.csv_upload(csv),self.staff); self.assertTrue(result["ok"]); self.assertTrue(Question.objects.filter(code="NEW-1",options__is_correct=True).exists())
    def test_invalid_multi_is_atomic(self):
        csv="question_code,module_code,category_code,category_title,question_type,question_text,option_a,option_b,option_c,option_d,option_e,option_f,correct_options,marks,difficulty,explanation,image_filename,image_alt_text,status\nBAD-1,MET,METAR,Weather,MULTIPLE,What?,Yes,No,,,,,A,1,EASY,,,,DRAFT\n"
        result=import_questions(self.csv_upload(csv),self.staff); self.assertFalse(result["ok"]); self.assertFalse(Question.objects.filter(code="BAD-1").exists())
    def test_zip_image_import_and_hash(self):
        im=Image.new("RGB",(30,20),(50,100,150)); ib=io.BytesIO(); im.save(ib,"PNG")
        csv="question_code,module_code,category_code,category_title,question_type,question_text,option_a,option_b,option_c,option_d,option_e,option_f,correct_options,marks,difficulty,explanation,image_filename,image_alt_text,status\nIMG-1,MET,CHART,Charts,SINGLE,Read image,One,Two,,,,,A,1,MEDIUM,,chart.png,Blue chart diagram,DRAFT\n"
        zb=io.BytesIO()
        with zipfile.ZipFile(zb,"w") as z: z.writestr("questions.csv",csv); z.writestr("images/chart.png",ib.getvalue())
        result=import_questions(SimpleUploadedFile("bank.zip",zb.getvalue()),self.staff); q=Question.objects.get(code="IMG-1"); self.assertTrue(result["ok"]); self.assertEqual(len(q.image_sha256),64); self.assertEqual(q.image_alt_text,"Blue chart diagram")
    def test_zip_path_traversal_rejected(self):
        zb=io.BytesIO()
        with zipfile.ZipFile(zb,"w") as z: z.writestr("questions.csv","x\n"); z.writestr("../evil.png",b"no")
        with self.assertRaises(ValueError): import_questions(SimpleUploadedFile("bad.zip",zb.getvalue()),self.staff)
    def test_duplicate_image_basenames_rejected(self):
        im=Image.new("RGB",(2,2)); b=io.BytesIO(); im.save(b,"PNG")
        zb=io.BytesIO()
        with zipfile.ZipFile(zb,"w") as z: z.writestr("questions.csv","x\n"); z.writestr("one/same.png",b.getvalue()); z.writestr("two/same.png",b.getvalue())
        with self.assertRaises(ValueError): import_questions(SimpleUploadedFile("bad.zip",zb.getvalue()),self.staff)


class SimpleImportTests(Base):
    SIMPLE_HEADER = "Module,Category,Question,Option A,Option B,Option C,Option D,Correct Answer,Image,Image Description\n"

    def simple_csv(self, row):
        return SimpleUploadedFile("simple.csv", (self.SIMPLE_HEADER + row + "\n").encode("utf-8"), content_type="text/csv")

    def test_simple_import_infers_type_and_defaults(self):
        upload = self.simple_csv("Meteorology,Clouds,Select cloud types,Cumulus,Stone,Cirrus,,A;C,,")
        result = import_simple_questions(upload, [], self.staff)
        self.assertTrue(result["ok"])
        question = Question.objects.get(stem="Select cloud types")
        self.assertEqual(question.question_type, Question.MULTIPLE)
        self.assertEqual(question.status, Question.DRAFT)
        self.assertEqual(question.marks, 1)
        self.assertEqual(set(question.options.filter(is_correct=True).values_list("key", flat=True)), {"A", "C"})
        self.assertTrue(question.code.startswith("METEOROLOGY-"))

    def test_simple_import_uses_general_category_and_single_answer(self):
        upload = self.simple_csv("Air Law,,Choose one,Yes,No,,,B,,")
        result = import_simple_questions(upload, [], self.staff)
        question = Question.objects.get(stem="Choose one")
        self.assertTrue(result["ok"])
        self.assertEqual(question.category.code, "GENERAL")
        self.assertEqual(question.question_type, Question.SINGLE)

    def test_simple_import_with_separately_selected_image(self):
        image = Image.new("RGB", (20, 20), (40, 80, 120))
        data = io.BytesIO(); image.save(data, "PNG")
        picture = SimpleUploadedFile("chart.png", data.getvalue(), content_type="image/png")
        upload = self.simple_csv("Navigation,Charts,Read the chart,North,South,,,A,chart.png,Blue navigation chart")
        result = import_simple_questions(upload, [picture], self.staff)
        question = Question.objects.get(stem="Read the chart")
        self.assertTrue(result["ok"])
        self.assertEqual(question.image_alt_text, "Blue navigation chart")
        self.assertEqual(len(question.image_sha256), 64)

    def test_simple_import_is_atomic_on_bad_row(self):
        csv_text = self.SIMPLE_HEADER + "Air Law,Rules,Good?,Yes,No,,,A,,\nAir Law,Rules,Bad?,Yes,No,,,Z,,\n"
        result = import_simple_questions(SimpleUploadedFile("simple.csv", csv_text.encode()), [], self.staff)
        self.assertFalse(result["ok"])
        self.assertFalse(Question.objects.filter(stem__in=["Good?", "Bad?"]).exists())

    def test_staff_import_page_accepts_one_step_upload(self):
        self.client.login(username="admin", password="Strong-pass-475")
        upload = self.simple_csv("Operations,General,Ready?,Yes,No,,,A,,")
        response = self.client.post(reverse("import_questions"), {"spreadsheet": upload})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "1 question imported")
        self.assertTrue(Question.objects.filter(stem="Ready?").exists())

    def test_staff_form_accepts_picture_without_zip(self):
        self.client.login(username="admin", password="Strong-pass-475")
        image = Image.new("RGB", (12, 12), (20, 60, 100))
        data = io.BytesIO(); image.save(data, "PNG")
        picture = SimpleUploadedFile("direct.png", data.getvalue(), content_type="image/png")
        upload = self.simple_csv("Navigation,Charts,Direct picture?,Yes,No,,,A,direct.png,Small blue test image")
        response = self.client.post(reverse("import_questions"), {"spreadsheet": upload, "images": [picture]})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "1 question imported")
        self.assertTrue(Question.objects.get(stem="Direct picture?").image.name.endswith(".png"))

    def test_simple_template_has_nine_columns(self):
        self.client.login(username="admin", password="Strong-pass-475")
        response = self.client.get(reverse("import_template"))
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        headers = [cell.value for cell in next(workbook["Questions"].iter_rows())]
        self.assertEqual(headers, ["Category", "Question", "Option A", "Option B", "Option C", "Option D", "Correct Answer", "Image", "Image Description"])

    def test_format_agnostic_image_matching(self):
        image = Image.new("RGB", (20, 20), (40, 80, 120))
        data = io.BytesIO(); image.save(data, "PNG")
        picture = SimpleUploadedFile("clouds.png", data.getvalue(), content_type="image/png")
        # In the CSV, we specify just 'clouds' without any extension
        upload = self.simple_csv("Air Law,Weather,Identify clouds,Cumulus,Stratus,,,A,clouds,Fluffy white clouds")
        result = import_simple_questions(upload, [picture], self.staff)
        self.assertTrue(result["ok"])
        q = Question.objects.get(stem="Identify clouds")
        self.assertTrue(q.image.name.endswith(".png"))
        self.assertEqual(q.image_alt_text, "Fluffy white clouds")

    def test_stratified_sampling_balances_sections(self):
        cat1 = Category.objects.create(module=self.module, code="AIRSPACE", title="Airspace")
        cat2 = Category.objects.create(module=self.module, code="REGS", title="Regulations")
        cat3 = Category.objects.create(module=self.module, code="ACCIDENTS", title="Accidents")
        Question.objects.filter(module=self.module).delete()

        for i in range(15):
            q = Question.objects.create(code=f"AIR_{i}", module=self.module, category=cat1, question_type=Question.SINGLE, stem=f"Airspace {i}", status=Question.PUBLISHED, created_by=self.staff)
            Option.objects.create(question=q, key="A", text="Yes", is_correct=True); Option.objects.create(question=q, key="B", text="No")
        for i in range(10):
            q = Question.objects.create(code=f"REG_{i}", module=self.module, category=cat2, question_type=Question.SINGLE, stem=f"Reg {i}", status=Question.PUBLISHED, created_by=self.staff)
            Option.objects.create(question=q, key="A", text="Yes", is_correct=True); Option.objects.create(question=q, key="B", text="No")
        for i in range(3):
            q = Question.objects.create(code=f"ACC_{i}", module=self.module, category=cat3, question_type=Question.SINGLE, stem=f"Accident {i}", status=Question.PUBLISHED, created_by=self.staff)
            Option.objects.create(question=q, key="A", text="Yes", is_correct=True); Option.objects.create(question=q, key="B", text="No")

        exam = Exam.objects.create(code="STRAT-EXAM", title="Stratified Test", module=self.module, duration_minutes=30, question_count=9, pass_mark=75, max_attempts=2, status=Exam.PUBLISHED)
        Assignment.objects.create(exam=exam, candidate=self.user)

        attempt = start_attempt(exam, self.user)
        self.assertEqual(attempt.attempt_questions.count(), 9)

        from collections import Counter
        drawn_categories = Counter(attempt.attempt_questions.values_list("category_code", flat=True))
        self.assertEqual(drawn_categories["ACCIDENTS"], 3)
        self.assertEqual(drawn_categories["REGS"], 3)
        self.assertEqual(drawn_categories["AIRSPACE"], 3)

    def test_stealth_mode_hides_future_exam_from_student_dashboard(self):
        # Schedule an exam in the future
        future_time = timezone.now() + timedelta(hours=2)
        future_exam = Exam.objects.create(
            code="FUTURE-EXAM",
            title="Future Secret Exam",
            module=self.module,
            duration_minutes=30,
            question_count=2,
            pass_mark=75,
            max_attempts=1,
            available_from=future_time,
            status=Exam.PUBLISHED,
        )
        Assignment.objects.create(exam=future_exam, candidate=self.user)

        self.client.login(username="C001", password="Strong-pass-473")
        response = self.client.get(reverse("dashboard"))
        # Candidate should NOT see the future exam in stealth mode
        self.assertNotContains(response, "Future Secret Exam")

        # Now test that when the time arrives (available_from in past), it becomes visible
        future_exam.available_from = timezone.now() - timedelta(minutes=5)
        future_exam.save()
        response2 = self.client.get(reverse("dashboard"))
        self.assertContains(response2, "Future Secret Exam")

    def test_palette_order_is_strictly_sequential(self):
        attempt = start_attempt(self.exam, self.user)
        # Answer question 2 first
        aq2 = attempt.attempt_questions.get(position=2)
        save_response(attempt, aq2, ["A"], self.user)
        # Fetch question view
        self.client.login(username="C001", password="Strong-pass-473")
        response = self.client.get(reverse("question", args=[attempt.id, 1]))
        palette = list(response.context["palette"])
        positions = [p["position"] for p in palette]
        self.assertEqual(positions, [1, 2, 3])

    def test_exam_disappears_when_max_attempts_reached(self):
        self.exam.max_attempts = 1
        self.exam.save()
        attempt = start_attempt(self.exam, self.user)
        submit_attempt(attempt, actor=self.user)

        self.client.login(username="C001", password="Strong-pass-473")
        response = self.client.get(reverse("dashboard"))
        # Active exam card must be gone
        self.assertEqual(len(response.context["rows"]), 0)
        self.assertContains(response, "No Examinations Currently Assigned")
        # Completed attempt appears in recent table
        self.assertEqual(len(response.context["recent"]), 1)

    def test_exam_disappears_when_available_until_passed(self):
        self.exam.available_until = timezone.now() - timedelta(minutes=10)
        self.exam.save()
        self.client.login(username="C001", password="Strong-pass-473")
        response = self.client.get(reverse("dashboard"))
        # Active exam card must be gone
        self.assertEqual(len(response.context["rows"]), 0)
        self.assertContains(response, "No Examinations Currently Assigned")

    def test_get_on_answer_endpoint_redirects_gracefully_to_question(self):
        attempt = start_attempt(self.exam, self.user)
        self.client.login(username="C001", password="Strong-pass-473")
        # Simulates candidate logging in with ?next pointing to an answer endpoint
        response = self.client.get(reverse("answer", args=[attempt.id, 2]))
        self.assertEqual(response.status_code, 302)
        self.assertIn(f"/attempts/{attempt.id}/questions/2/", response.url)

    def test_staff_can_view_student_attempt_question(self):
        attempt = start_attempt(self.exam, self.user)
        # Login as staff (admin)
        self.client.login(username="admin", password="Strong-pass-475")
        response = self.client.get(reverse("question", args=[attempt.id, 1]))
        self.assertEqual(response.status_code, 200)

    def test_confirm_submit_shows_flagged_and_unanswered_question_numbers(self):
        attempt = start_attempt(self.exam, self.user)
        # Flag question 2
        aq2 = attempt.attempt_questions.get(position=2)
        aq2.flagged = True
        aq2.save()
        # Answer question 1
        aq1 = attempt.attempt_questions.get(position=1)
        save_response(attempt, aq1, ["A"], self.user)

        self.client.login(username="C001", password="Strong-pass-473")
        response = self.client.get(reverse("confirm_submit", args=[attempt.id]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["flagged_positions"], [2])
        self.assertIn(2, response.context["unanswered_positions"])
        self.assertIn(3, response.context["unanswered_positions"])
